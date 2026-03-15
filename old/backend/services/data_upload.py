from io import BytesIO

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from repo.data_upload import DataUploadRepo


class DataUploadService:
    def __init__(self, repo: DataUploadRepo | None = None):
        self.repo = repo or DataUploadRepo()

    async def upload_file(self, session: Session, file: UploadFile) -> dict:
        filename = file.filename or ""
        ext = filename.split(".")[-1].lower() if "." in filename else ""

        if ext not in {"xlsx", "pdf"}:
            raise HTTPException(status_code=400, detail="Unsupported file format. Use XLSX or PDF.")

        if ext == "pdf":
            return {
                "status": "accepted",
                "fileType": "pdf",
                "uploaded": 0,
                "skipped": 0,
                "message": "PDF upload is accepted.",
            }

        content = await file.read()
        workbook = load_workbook(filename=BytesIO(content), data_only=True)
        sheet = workbook.active

        headers = [str(cell.value).strip().lower() if cell.value is not None else "" for cell in sheet[1]]
        header_index = {name: idx for idx, name in enumerate(headers)}

        required = {"indicator", "region", "year", "value"}
        if not required.issubset(set(header_index)):
            raise HTTPException(
                status_code=400,
                detail="XLSX must contain headers: indicator, region, year, value.",
            )

        source_col = "source" if "source" in header_index else None
        source_id_col = "source_id" if "source_id" in header_index else None

        rows_to_save: list[dict] = []
        skipped = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            indicator_name = row[header_index["indicator"]]
            region_name = row[header_index["region"]]
            year = row[header_index["year"]]
            value = row[header_index["value"]]

            if indicator_name is None or region_name is None or year is None or value is None:
                skipped += 1
                continue

            indicator_id = self.repo.find_indicator_id(session, str(indicator_name).strip())
            region_id = self.repo.find_region_id(session, str(region_name).strip())

            if not indicator_id or not region_id:
                skipped += 1
                continue

            source_id = None
            if source_id_col:
                raw_source_id = row[header_index[source_id_col]]
                source_id = int(raw_source_id) if raw_source_id is not None else None
            elif source_col:
                source_name = row[header_index[source_col]]
                source_id = self.repo.find_source_id(session, str(source_name).strip()) if source_name else None

            if source_id is None:
                source_id = 1

            rows_to_save.append(
                {
                    "indicator_id": int(indicator_id),
                    "region_id": int(region_id),
                    "year": int(year),
                    "value": float(value),
                    "source_id": int(source_id),
                }
            )

        uploaded = self.repo.save_indicator_values(session=session, rows=rows_to_save) if rows_to_save else 0
        return {
            "status": "success",
            "fileType": "xlsx",
            "uploaded": uploaded,
            "skipped": skipped,
        }
